"""Scrap statistics over a date range (Data log > Scrap statistics).

Parts are counted the same way as the camera view's OK/NOK chart: each
reading carries the running totals of its job, so the parts made between two
consecutive readings of the same job are the difference of their totals, and
a job change starts a new baseline. Scrap is fail / (pass + fail).

Days are calendar days in the viewer's time zone; the range is inclusive.
"""
from __future__ import annotations

import datetime as dt
import io
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from sqlalchemy.orm import Session

from .models import Device, Reading


def zone(name: str | None) -> dt.tzinfo:
    try:
        return ZoneInfo(name) if name else dt.timezone.utc
    except (ZoneInfoNotFoundError, ValueError):
        return dt.timezone.utc


def _aware(t: dt.datetime) -> dt.datetime:
    return t.replace(tzinfo=dt.timezone.utc) if t.tzinfo is None else t


def _row(**keys) -> dict:
    return {**keys, "pass": 0, "fail": 0}


def _finish(rows: list[dict]) -> list[dict]:
    for r in rows:
        r["total"] = r["pass"] + r["fail"]
        r["scrap_rate"] = round(r["fail"] / r["total"], 4) if r["total"] else 0.0
    return rows


def compute(db: Session, first_day: dt.date, last_day: dt.date, tz: dt.tzinfo) -> dict:
    start = dt.datetime.combine(first_day, dt.time(), tz).astimezone(dt.timezone.utc)
    end = dt.datetime.combine(last_day + dt.timedelta(days=1), dt.time(), tz).astimezone(dt.timezone.utc)

    days = {}
    d = first_day
    while d <= last_day:
        days[d] = _row(day=d.isoformat())
        d += dt.timedelta(days=1)
    overall = _row()
    per_camera: dict[int, dict] = {}
    per_job: dict[tuple[int, str], dict] = {}

    cols = (Reading.job_name, Reading.total_pass, Reading.total_fail, Reading.created_at)
    for dev in db.query(Device).order_by(Device.name).all():
        cam = per_camera[dev.id] = _row(camera_id=dev.id, camera=dev.name)
        q = db.query(*cols).filter(Reading.device_id == dev.id)
        baseline = q.filter(Reading.created_at < start).order_by(Reading.created_at.desc(), Reading.id.desc()).first()
        rows = (
            q.filter(Reading.created_at >= start, Reading.created_at < end)
            .order_by(Reading.created_at.asc(), Reading.id.asc())
            .yield_per(5000)
        )
        prev = baseline
        for r in rows:
            if prev is not None and prev.job_name == r.job_name:
                d_ok = max(r.total_pass - prev.total_pass, 0)
                d_nok = max(r.total_fail - prev.total_fail, 0)
                if d_ok or d_nok:
                    day = _aware(r.created_at).astimezone(tz).date()
                    job = r.job_name or "—"
                    jrow = per_job.setdefault((dev.id, job), _row(camera=dev.name, job=job))
                    for target in (overall, cam, jrow, days[day]):
                        target["pass"] += d_ok
                        target["fail"] += d_nok
            prev = r

    return {
        "from": first_day.isoformat(),
        "to": last_day.isoformat(),
        "overall": _finish([overall])[0],
        "per_camera": _finish(list(per_camera.values())),
        "per_job": _finish(sorted(per_job.values(), key=lambda r: (r["camera"], r["job"]))),
        "per_day": _finish(list(days.values())),
    }


def to_xlsx(stats: dict) -> bytes:
    """One sheet per view, with the same figures as the screen."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    period = f"{stats['from']} to {stats['to']}"
    sheets = [
        ("Overall", [("Period", None)], [{**stats["overall"], "period": period}], ["period"]),
        ("Per camera", [("Camera", None)], stats["per_camera"], ["camera"]),
        ("Per job", [("Camera", None), ("Job", None)], stats["per_job"], ["camera", "job"]),
        ("Per day", [("Day", None)], stats["per_day"], ["day"]),
    ]
    for i, (title, heads, rows, keys) in enumerate(sheets):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = title
        ws.append([h for h, _ in heads] + ["Pass", "Fail", "Total", "Scrap %"])
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            day = [dt.date.fromisoformat(r["day"])] if keys == ["day"] else []
            ws.append(day + [r[k] for k in keys if k != "day"] + [r["pass"], r["fail"], r["total"], r["scrap_rate"]])
        n = len(keys)
        for row in ws.iter_rows(min_row=2):
            row[n + 3].number_format = "0.00%"
            if keys == ["day"]:
                row[0].number_format = "yyyy-mm-dd"
        for col, width in zip("ABCDEFG", [24] * n + [12, 12, 12, 10]):
            ws.column_dimensions[col].width = width
        ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
